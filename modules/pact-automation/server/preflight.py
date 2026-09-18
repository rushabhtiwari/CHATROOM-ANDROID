"""
server/preflight.py - parse an uploaded CSV/XLSX and check every row BEFORE anything runs.

Nothing here touches the GUI. It answers one question: if we hand these rows to the robot,
which ones will blow up? Blocking problems are `errors` (the row can be excluded with a
checkbox on the page); everything else is a `warning` and is shown, not blocked.

    rows = parse(filename, data)                 # -> [{"row_no": 1, "record": {...}}, ...]
    checked = validate(rows, profile)            # -> rows + "errors" / "warnings"
"""
import csv, io, re
from datetime import datetime
from pathlib import Path

# field-name heuristics, used when a profile field carries no explicit "validate"
_NUMBER_HINT = re.compile(r"(qty|quantity|amount|amt|price|rate|limit|total|value|discount)", re.I)
_EMAIL_HINT = re.compile(r"e?mail", re.I)
_PHONE_HINT = re.compile(r"(phone|mobile|contact_no|telephone)", re.I)
_DATE_HINT = re.compile(r"(date|validity)", re.I)

_NUMBER_RE = re.compile(r"^-?[\d,]*\.?\d+$")
_DATE_RE = re.compile(r"^\d{1,2}[/-]\d{1,2}[/-]\d{2,4}$")

LINE_ITEMS_COLUMN = "line_items"


# --------------------------------------------------------------------------- parsing
def parse(filename: str, data: bytes) -> list[dict]:
    """CSV or XLSX bytes -> [{"row_no": n, "record": {col: value}}]. Blank rows are dropped."""
    ext = Path(filename or "").suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        table = _parse_xlsx(data)
    else:
        table = _parse_csv(data)
    out = []
    for i, raw in enumerate(table, start=1):
        rec = {(k or "").strip(): _clean(v) for k, v in raw.items() if (k or "").strip()}
        if not any(str(v).strip() for v in rec.values()):
            continue
        out.append({"row_no": i, "record": rec})
    return out


def _clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _parse_csv(data: bytes) -> list[dict]:
    text = data.decode("utf-8-sig", errors="replace")
    return list(csv.DictReader(io.StringIO(text)))


def _parse_xlsx(data: bytes) -> list[dict]:
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover - depends on the venv
        raise RuntimeError("openpyxl is not installed - run scripts\\start.ps1 to install it") from e
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    try:
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
    except StopIteration:
        return []
    out = []
    for r in rows:
        out.append({header[i]: (r[i] if i < len(r) else None) for i in range(len(header))})
    wb.close()
    return out


# --------------------------------------------------------------------------- line items
def parse_line_items(raw: str) -> tuple[list[dict], list[str]]:
    """`CODE|QTY|PRICE ; CODE|QTY|PRICE` -> [{product_code, qty, unit_price}], plus problems."""
    items, problems = [], []
    for chunk in str(raw or "").split(";"):
        chunk = chunk.strip()
        if not chunk:
            continue
        parts = [p.strip() for p in chunk.split("|")]
        if len(parts) != 3:
            problems.append(f"line item {chunk!r} is not CODE|QTY|PRICE")
            continue
        code, qty, price = parts
        if not code:
            problems.append(f"line item {chunk!r} has no product code")
        if not _NUMBER_RE.match(qty.replace(",", "")):
            problems.append(f"line item {code!r}: qty {qty!r} is not a number")
        if not _NUMBER_RE.match(price.replace(",", "")):
            problems.append(f"line item {code!r}: unit price {price!r} is not a number")
        items.append({"product_code": code, "qty": qty, "unit_price": price})
    return items, problems


def format_line_items(items: list[dict]) -> str:
    """The line items as the VERIFIER should judge them: code and quantity, no price.

    `UnitPrice` cannot be typed on the Purchase Order screen at all - the cell exposes no
    editor and PACT sources the rate from its own product master (the evidence is in
    demo/discovered_values.md). Telling the vision check to expect the customer's rate
    therefore asks it to fail every row whose master rate differs, which is not a data-entry
    error and is not something the robot could have done differently. It failed the whole
    document and blocked the save.

    What the customer's document said the rate should be is not lost: `Robot._check_price`
    reads back what PACT actually charged and records the difference as a price note on the
    entry, which is where a human can act on it.
    """
    return "; ".join(f"{i.get('product_code','')}|{i.get('qty','')}" for i in items)


# --------------------------------------------------------------------------- validation
def _kind(key: str, spec: dict) -> str:
    explicit = (spec or {}).get("validate")
    if explicit:
        return explicit
    t = (spec or {}).get("type", "edit")
    if t == "date":
        return "date"
    if _EMAIL_HINT.search(key):
        return "email"
    if _PHONE_HINT.search(key):
        return "phone"
    if _NUMBER_HINT.search(key):
        return "number"
    if _DATE_HINT.search(key):
        return "date"
    return "text"


_DATE_FORMATS = ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%m-%y")


def _as_date(value: str):
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    return None


def accounting_period(profile: dict):
    """(from, to) of the period PACT will post into, or None if the profile does not say."""
    ap = profile.get("accounting_period") or {}
    lo, hi = _as_date(ap.get("from", "")), _as_date(ap.get("to", ""))
    return (lo, hi) if lo and hi else None


def duplicate_key_fields(profile: dict) -> list[str]:
    keys = profile.get("duplicate_key")
    if keys:
        return list(keys)
    req = [k for k, v in profile.get("fields", {}).items() if v.get("required")]
    return req[:1] or list(profile.get("fields", {}))[:1]


def validate(rows: list[dict], profile: dict) -> list[dict]:
    fields = profile.get("fields", {})
    labels = profile.get("verify_labels", {})
    dup_fields = duplicate_key_fields(profile)
    grid = profile.get("grid")
    period = accounting_period(profile)
    seen: dict[str, int] = {}
    out = []

    for row in rows:
        rec = dict(row["record"])
        errors, warnings = [], []
        label = lambda k: labels.get(k, k)

        unknown = [k for k in rec if k not in fields and k != LINE_ITEMS_COLUMN]
        if unknown:
            warnings.append("column(s) not in the profile, will be ignored: " + ", ".join(unknown))

        # line items -> structured list on the record
        if LINE_ITEMS_COLUMN in rec:
            if not grid:
                warnings.append("this profile has no grid; line_items will be ignored")
                rec.pop(LINE_ITEMS_COLUMN)
            else:
                items, problems = parse_line_items(rec.pop(LINE_ITEMS_COLUMN))
                errors.extend(problems)
                if not items:
                    errors.append("line_items is empty - a purchase order needs at least one line")
                rec["line_items"] = items
        elif grid and grid.get("required", True):
            warnings.append("no line_items column - the document will be saved with an empty grid")

        for key, spec in fields.items():
            value = str(rec.get(key, "") or "").strip()
            if spec.get("required") and not value:
                errors.append(f"{label(key)} is required but empty")
                continue
            if not value:
                continue
            kind = _kind(key, spec)
            if kind == "number" and not _NUMBER_RE.match(value.replace(",", "")):
                errors.append(f"{label(key)} must be a number, got {value!r}")
            elif kind == "email" and "@" not in value:
                errors.append(f"{label(key)} is not an email address: {value!r}")
            elif kind == "phone" and re.search(r"[A-Za-z]", value):
                errors.append(f"{label(key)} contains letters: {value!r}")
            elif kind == "date" and not _DATE_RE.match(value):
                errors.append(f"{label(key)} must look like dd/MM/yyyy, got {value!r}")
            elif kind == "date" and period:
                d = _as_date(value)
                lo, hi = period
                if d is None:
                    errors.append(f"{label(key)} is not a real date: {value!r}")
                elif not (lo <= d <= hi):
                    errors.append(
                        f"{label(key)} {value} is outside the accounting period "
                        f"{lo.strftime('%d/%m/%Y')} to {hi.strftime('%d/%m/%Y')} - PACT will "
                        "refuse to post it")

            options = spec.get("options")
            if options and value not in options:
                lowered = {str(o).lower(): o for o in options}
                if value.lower() in lowered:
                    rec[key] = lowered[value.lower()]      # fix the casing, do not block
                    warnings.append(f"{label(key)}: corrected to {lowered[value.lower()]!r}")
                else:
                    errors.append(f"{label(key)}: {value!r} is not one of " + ", ".join(map(str, options)))
            elif spec.get("type") == "lookup_combo" and not options:
                warnings.append(f"{label(key)} is looked up in PACT's master - "
                                f"the row fails if {value!r} does not exist there")

        # duplicate key inside this same file
        dup_val = " | ".join(str(rec.get(k, "")).strip().lower() for k in dup_fields)
        if dup_val.strip(" |"):
            if dup_val in seen:
                errors.append(f"duplicate of row {seen[dup_val]} on "
                              + ", ".join(label(k) for k in dup_fields))
            else:
                seen[dup_val] = row["row_no"]

        out.append({"row_no": row["row_no"], "record": rec,
                    "errors": errors, "warnings": warnings, "include": not errors})
    return out
